from pathlib import Path
from docx import Document as DocxDocument
from backend.ingestion.normalizers.schema import ExtractedItem


class DocxParser:
    """
    Step 3 Parser: Extracts headings, paragraphs, and tables from DOCX files.
    """

    @classmethod
    def extract(cls, file_path: Path) -> list[ExtractedItem]:
        items: list[ExtractedItem] = []
        doc = DocxDocument(file_path)

        for idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""
            is_heading = "heading" in style_name or "title" in style_name

            items.append(
                ExtractedItem(
                    type="heading" if is_heading else "text",
                    text=text,
                    source={"paragraph_index": idx, "style": p.style.name if p.style else None},
                )
            )

        for t_idx, table in enumerate(doc.tables):
            rows_data: list[list[str]] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                cleaned_cells: list[str] = []
                for cell_val in cells:
                    cleaned_cells.append(cell_val)
                rows_data.append(cleaned_cells)

            if len(rows_data) > 1:
                headers = rows_data[0]
                rows = rows_data[1:]
                items.append(
                    ExtractedItem(
                        type="table",
                        headers=headers,
                        rows=rows,
                        source={"table_index": t_idx},
                    )
                )

        return items
